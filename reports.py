import os
from datetime import date, datetime
from io import BytesIO
import matplotlib.pyplot as plt
import matplotlib.dates as mdates

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

import database as db

plt.rcParams['font.family'] = 'DejaVu Sans'

FONT_REGISTERED = False


def register_fonts():
    global FONT_REGISTERED
    if FONT_REGISTERED:
        return

    font_paths = [
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        '/usr/share/fonts/TTF/DejaVuSans.ttf',
        '/usr/share/fonts/dejavu/DejaVuSans.ttf',
        'C:/Windows/Fonts/arial.ttf',
        '/System/Library/Fonts/Supplemental/Arial.ttf',
    ]

    for path in font_paths:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont('CustomFont', path))
                FONT_REGISTERED = True
                return
            except:
                continue


register_fonts()


# ==================== ГРАФИКИ ====================

def generate_chart(start_date: date = None, end_date: date = None,
                   chart_type: str = 'daily', rate: float = 100.0) -> BytesIO:
    if chart_type == 'daily':
        data = db.get_combined_daily_stats(start_date, end_date, rate)
        x_label, title = 'Дата', 'Доходы по дням'
    else:
        data = db.get_combined_monthly_stats(rate)
        x_label, title = 'Месяц', 'Доходы по месяцам'

    if not data:
        fig, ax = plt.subplots(figsize=(10, 6))
        ax.text(0.5, 0.5, 'Нет данных за выбранный период', ha='center', va='center', fontsize=14)
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
    else:
        fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(12, 10))

        if chart_type == 'daily':
            dates = [datetime.strptime(d['date'], '%Y-%m-%d') for d in data]
            rub_values = [d['total_rub'] for d in data]
            usd_values = [d['total_usd'] for d in data]

            ax1.fill_between(dates, rub_values, alpha=0.3, color='#2ecc71')
            ax1.plot(dates, rub_values, color='#2ecc71', linewidth=2, marker='o', markersize=4)
            ax1.xaxis.set_major_formatter(mdates.DateFormatter('%d.%m'))
            ax1.xaxis.set_major_locator(mdates.AutoDateLocator())
            ax1.set_title('В рублях (₽)', fontsize=12, fontweight='bold')
            ax1.set_ylabel('Сумма (₽)')
            ax1.grid(True, alpha=0.3)
            ax1.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, p: f'{x:,.0f}'))
            plt.setp(ax1.xaxis.get_majorticklabels(), rotation=45)

            ax2.fill_between(dates, usd_values, alpha=0.3, color='#3498db')
            ax2.plot(dates, usd_values, color='#3498db', linewidth=2, marker='o', markersize=4)
            ax2.xaxis.set_major_formatter(mdates.DateFormatter('%d.%m'))
            ax2.xaxis.set_major_locator(mdates.AutoDateLocator())
            ax2.set_title('В долларах ($)', fontsize=12, fontweight='bold')
            ax2.set_xlabel(x_label)
            ax2.set_ylabel('Сумма ($)')
            ax2.grid(True, alpha=0.3)
            ax2.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, p: f'{x:,.0f}'))
            plt.setp(ax2.xaxis.get_majorticklabels(), rotation=45)
        else:
            months = [d['month'] for d in data]
            rub_values = [d['total_rub'] for d in data]
            usd_values = [d['total_usd'] for d in data]

            bars1 = ax1.bar(months, rub_values, color='#2ecc71', alpha=0.8)
            ax1.set_title('В рублях (₽)', fontsize=12, fontweight='bold')
            ax1.set_ylabel('Сумма (₽)')
            for bar, val in zip(bars1, rub_values):
                if val > 0:
                    ax1.annotate(f'{val:,.0f}', xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()),
                                 xytext=(0, 3), textcoords="offset points", ha='center', va='bottom', fontsize=8)
            plt.setp(ax1.xaxis.get_majorticklabels(), rotation=45)
            ax1.grid(True, alpha=0.3, axis='y')

            bars2 = ax2.bar(months, usd_values, color='#3498db', alpha=0.8)
            ax2.set_title('В долларах ($)', fontsize=12, fontweight='bold')
            ax2.set_xlabel(x_label)
            ax2.set_ylabel('Сумма ($)')
            for bar, val in zip(bars2, usd_values):
                if val > 0:
                    ax2.annotate(f'{val:,.0f}', xy=(bar.get_x() + bar.get_width() / 2, bar.get_height()),
                                 xytext=(0, 3), textcoords="offset points", ha='center', va='bottom', fontsize=8)
            plt.setp(ax2.xaxis.get_majorticklabels(), rotation=45)
            ax2.grid(True, alpha=0.3, axis='y')

        fig.suptitle(f'{title}\n(курс: 1$ = {rate:.2f}₽)', fontsize=14, fontweight='bold')

    plt.tight_layout()
    buf = BytesIO()
    plt.savefig(buf, format='png', dpi=150, bbox_inches='tight', facecolor='white', edgecolor='none')
    buf.seek(0)
    plt.close(fig)
    return buf


# ==================== EXCEL - ИСПРАВЛЕНО ====================

def set_column_widths(ws, widths: dict):
    """Установка ширины колонок по словарю {номер_колонки: ширина}"""
    for col_num, width in widths.items():
        ws.column_dimensions[get_column_letter(col_num)].width = width


def generate_excel_report(start_date: date = None, end_date: date = None,
                          rate: float = 100.0) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"

    # Стили
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    title_font = Font(bold=True, size=14)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    num_fmt = '#,##0.00'

    # Заголовок
    ws['A1'] = "Отчёт по продажам"
    ws['A1'].font = title_font

    period_text = "За всё время"
    if start_date and end_date:
        period_text = f"{start_date.strftime('%d.%m.%Y')} - {end_date.strftime('%d.%m.%Y')}"

    ws['A2'] = f"Период: {period_text} | Курс: 1$ = {rate:.2f} RUB"
    ws['A3'] = f"Сформирован: {datetime.now().strftime('%d.%m.%Y %H:%M')}"

    # Статистика
    stats = db.get_combined_stats(start_date, end_date, rate)

    ws['A5'] = "СВОДКА"
    ws['A5'].font = Font(bold=True)

    summary = [
        ("Сделок:", stats['total_count'], ""),
        ("Итого RUB:", stats['total_rub'], num_fmt),
        ("Итого USD:", stats['total_usd'], num_fmt),
        ("Средний чек RUB:", stats['avg_rub'], num_fmt),
        ("Средний чек USD:", stats['avg_usd'], num_fmt),
    ]

    for i, (label, value, fmt) in enumerate(summary):
        ws.cell(row=6 + i, column=1, value=label)
        cell = ws.cell(row=6 + i, column=2, value=value)
        if fmt:
            cell.number_format = fmt

    # Таблица транзакций
    start_row = 6 + len(summary) + 2
    ws.cell(row=start_row, column=1, value="ТРАНЗАКЦИИ").font = Font(bold=True)

    headers = ['ID', 'Дата', 'Сумма', 'Валюта', 'В RUB', 'В USD', 'Комментарий']
    header_row = start_row + 1

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    transactions = db.get_all_transactions(start_date, end_date)

    for i, t in enumerate(transactions):
        row = header_row + 1 + i

        if t['currency'] == 'RUB':
            in_rub, in_usd = t['amount'], t['amount'] / rate if rate > 0 else 0
        else:
            in_usd, in_rub = t['amount'], t['amount'] * rate

        values = [t['id'], t['transaction_date'], t['amount'], t['currency'], in_rub, in_usd, t['comment'] or '-']

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = border
            if col in [3, 5, 6]:  # Числовые колонки
                cell.number_format = num_fmt

    # Ширина колонок (фиксированная, без итерации)
    set_column_widths(ws, {1: 8, 2: 12, 3: 15, 4: 10, 5: 15, 6: 15, 7: 30})

    # Лист по дням
    ws2 = wb.create_sheet("По дням")
    daily_stats = db.get_combined_daily_stats(start_date, end_date, rate)

    headers2 = ['Дата', 'В RUB', 'В USD', 'Сделок']
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border

    for i, d in enumerate(daily_stats):
        row = 2 + i
        ws2.cell(row=row, column=1, value=d['date']).border = border

        cell_rub = ws2.cell(row=row, column=2, value=d['total_rub'])
        cell_rub.border = border
        cell_rub.number_format = num_fmt

        cell_usd = ws2.cell(row=row, column=3, value=d['total_usd'])
        cell_usd.border = border
        cell_usd.number_format = num_fmt

        ws2.cell(row=row, column=4, value=d['count']).border = border

    set_column_widths(ws2, {1: 12, 2: 18, 3: 18, 4: 10})

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ==================== PDF ====================

def generate_pdf_report(start_date: date = None, end_date: date = None,
                        rate: float = 100.0) -> BytesIO:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            rightMargin=20 * mm, leftMargin=20 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm)

    font_name = 'CustomFont' if FONT_REGISTERED else 'Helvetica'

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', fontName=font_name, fontSize=18, spaceAfter=10, alignment=1)
    heading_style = ParagraphStyle('Heading', fontName=font_name, fontSize=14, spaceBefore=15, spaceAfter=10)
    normal_style = ParagraphStyle('Normal', fontName=font_name, fontSize=10)
    center_style = ParagraphStyle('Center', fontName=font_name, fontSize=10, alignment=1)

    elements = []

    elements.append(Paragraph("Отчёт по продажам", title_style))

    period_text = "За всё время"
    if start_date and end_date:
        period_text = f"{start_date.strftime('%d.%m.%Y')} - {end_date.strftime('%d.%m.%Y')}"

    elements.append(Paragraph(f"Период: {period_text}", center_style))
    elements.append(Paragraph(f"Курс: 1$ = {rate:.2f} RUB", center_style))
    elements.append(Spacer(1, 20))

    stats = db.get_combined_stats(start_date, end_date, rate)

    elements.append(Paragraph("Сводная статистика", heading_style))

    stats_data = [
        ["Показатель", "Значение"],
        ["Всего сделок", str(stats['total_count'])],
        ["Итого RUB", f"{stats['total_rub']:,.2f}"],
        ["Итого USD", f"{stats['total_usd']:,.2f}"],
        ["Средний чек RUB", f"{stats['avg_rub']:,.2f}"],
        ["Средний чек USD", f"{stats['avg_usd']:,.2f}"],
    ]

    stats_table = Table(stats_data, colWidths=[80 * mm, 80 * mm])
    stats_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2E7D32')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#81C784')),
    ]))
    elements.append(stats_table)
    elements.append(Spacer(1, 20))

    # График
    try:
        chart_buf = generate_chart(start_date, end_date, 'daily', rate)
        chart_img = Image(chart_buf, width=160 * mm, height=100 * mm)
        elements.append(Paragraph("График доходов", heading_style))
        elements.append(chart_img)
        elements.append(Spacer(1, 20))
    except:
        pass

    # Транзакции
    elements.append(Paragraph("Транзакции", heading_style))

    transactions = db.get_all_transactions(start_date, end_date)

    if transactions:
        trans_data = [["Дата", "Сумма", "RUB", "USD"]]
        for t in transactions[:50]:
            if t['currency'] == 'RUB':
                in_rub, in_usd = t['amount'], t['amount'] / rate if rate > 0 else 0
                amt = f"{t['amount']:,.0f} RUB"
            else:
                in_usd, in_rub = t['amount'], t['amount'] * rate
                amt = f"{t['amount']:,.0f} USD"

            trans_data.append([t['transaction_date'], amt, f"{in_rub:,.0f}", f"{in_usd:,.0f}"])

        trans_table = Table(trans_data, colWidths=[35 * mm, 40 * mm, 40 * mm, 40 * mm])
        trans_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2E7D32')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, -1), font_name),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        elements.append(trans_table)

        if len(transactions) > 50:
            elements.append(Paragraph(f"...и ещё {len(transactions) - 50} транзакций", normal_style))

    doc.build(elements)
    buf.seek(0)
    return buf